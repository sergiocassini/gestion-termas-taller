import streamlit as st
import pandas as pd
import psycopg2
from datetime import datetime, date
import base64
import io
import openpyxl
from openpyxl.drawing.image import Image as OpenPyXLEImage
from PIL import Image as PILImage

# Configuración de página adaptable a teléfonos
st.set_page_config(page_title="Gestión Termas & Taller", page_icon="🏪", layout="wide")

# --- CONTROL DE ACCESO CON PIN ---
PIN_CORRECTO = "2017"

if "autenticado" not in st.session_state:
    st.session_state["autenticado"] = False

def verificar_pin():
    if str(st.session_state["input_pin"]).strip() == PIN_CORRECTO:
        st.session_state["autenticado"] = True
    else:
        st.error("🔒 PIN / Contraseña incorrecta")

if not st.session_state["autenticado"]:
    st.title("🔒 Acceso Restringido")
    st.subheader("Control de Negocio: Termas, Taller & Feria")
    st.text_input("Ingresá el PIN de acceso:", type="password", key="input_pin", on_change=verificar_pin)
    st.button("Ingresar", on_click=verificar_pin, use_container_width=True)
    st.stop()

# --- CONEXIÓN A SUPABASE ---
def get_connection():
    return psycopg2.connect(
        host=st.secrets["postgres"]["host"],
        database=st.secrets["postgres"]["database"],
        user=st.secrets["postgres"]["user"],
        password=st.secrets["postgres"]["password"],
        port=st.secrets["postgres"]["port"]
    )

def consulta(query, params=(), fetch=True):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(query, params)
    conn.commit()
    res = cursor.fetchall() if fetch else None
    conn.close()
    return res

# --- HELPER PARA PROCESAR IMÁGENES ---
def procesar_archivo_imagen(archivo_subido):
    if archivo_subido is None:
        return None
    bytes_data = archivo_subido.getvalue()
    mime_type = archivo_subido.type if archivo_subido.type else "image/png"
    b64_str = base64.b64encode(bytes_data).decode('utf-8')
    return f"data:{mime_type};base64,{b64_str}"

# --- INICIALIZACIÓN Y AUTO-MIGRACIÓN ---
def init_db():
    try:
        consulta("ALTER TABLE productos ADD COLUMN IF NOT EXISTS imagen_url TEXT;", fetch=False)
        consulta("ALTER TABLE productos ADD COLUMN IF NOT EXISTS grupo VARCHAR(100);", fetch=False)
        consulta("ALTER TABLE productos ADD COLUMN IF NOT EXISTS subgrupo VARCHAR(100) DEFAULT 'Varios';", fetch=False)
        consulta("ALTER TABLE gachapon_premios ADD COLUMN IF NOT EXISTS coleccion VARCHAR(100) DEFAULT 'Zooki';", fetch=False)
        consulta("ALTER TABLE gachapon_premios ADD COLUMN IF NOT EXISTS precio_ficha NUMERIC DEFAULT 2000;", fetch=False)
        consulta("ALTER TABLE ventas ADD COLUMN IF NOT EXISTS nombre_feria VARCHAR(150) DEFAULT 'General';", fetch=False)
        consulta("ALTER TABLE ventas ADD COLUMN IF NOT EXISTS id_cierre INT DEFAULT 0;", fetch=False)
    except Exception:
        pass

    consulta("""
        CREATE TABLE IF NOT EXISTS productos (
            id SERIAL PRIMARY KEY,
            nombre VARCHAR(255) UNIQUE NOT NULL,
            categoria VARCHAR(100) DEFAULT 'General',
            grupo VARCHAR(100),
            subgrupo VARCHAR(100) DEFAULT 'Varios',
            precio NUMERIC DEFAULT 0,
            stock INT DEFAULT 0,
            stock_minimo INT DEFAULT 1,
            imagen_url TEXT
        );
        CREATE TABLE IF NOT EXISTS gachapon_premios (
            id SERIAL PRIMARY KEY,
            numero VARCHAR(50),
            nombre VARCHAR(255) UNIQUE NOT NULL,
            coleccion VARCHAR(100) DEFAULT 'Zooki',
            impresos INT DEFAULT 0,
            stock_deposito INT DEFAULT 0,
            en_maquina INT DEFAULT 0,
            precio_ficha NUMERIC DEFAULT 2000,
            imagen_url TEXT
        );
        CREATE TABLE IF NOT EXISTS ventas (
            id SERIAL PRIMARY KEY,
            origen VARCHAR(50) DEFAULT 'TERMAS',
            nombre_feria VARCHAR(150) DEFAULT 'General',
            item_tipo VARCHAR(50) NOT NULL,
            item_nombre VARCHAR(255) NOT NULL,
            cantidad INT NOT NULL,
            precio_unitario NUMERIC NOT NULL,
            subtotal NUMERIC NOT NULL,
            fecha TIMESTAMP NOT NULL,
            cerrado INT DEFAULT 0,
            id_cierre INT DEFAULT 0
        );
        CREATE TABLE IF NOT EXISTS cierres_caja (
            id SERIAL PRIMARY KEY,
            tipo_cierre VARCHAR(50) NOT NULL,
            etiqueta_ciclo VARCHAR(150) NOT NULL,
            total_recaudado NUMERIC NOT NULL,
            total_operaciones INT NOT NULL,
            fecha_inicio TIMESTAMP,
            fecha_cierre TIMESTAMP NOT NULL
        );
    """, fetch=False)

try:
    init_db()
except Exception as e:
    st.error(f"Conectando con la base de datos... ({e})")

# --- MAPEO INTELIGENTE DE GRUPOS ---
def obtener_grupo(categoria, nombre=""):
    text = (str(categoria) + " " + str(nombre)).lower()
    if any(k in text for k in ['llavero', 'dije', 'pin', 'iman', 'dijes', 'set imanes']):
        return "🔑 Llaveros, Pines y Dijes"
    elif any(k in text for k in ['fidget', 'juguete', 'articulado', 'kit cards', 'burbujero']):
        return "🧩 Juguetes y Fidgets"
    elif any(k in text for k in ['figura', 'funko', 'colgante']):
        return "🗿 Figuras y Funkos"
    elif any(k in text for k in ['mate', 'vaso', 'maceta', 'portacelular', 'sahumerios', 'porta velas', 'floreros']):
        return "🏠 Hogar y Deco"
    else:
        return "📦 Varios y Novedades"

# --- RENDERIZADOR DE CATÁLOGO VISUAL ---
def renderizar_catalogo(prods, origen_venta, incluir_gachapon=False, key_prefix="cat", nombre_feria="General"):
    busqueda = st.text_input("🔍 Buscador rápido de producto (nombre o categoría):", key=f"busqueda_{key_prefix}")
    
    if prods:
        df_prods = pd.DataFrame(prods, columns=["id", "nombre", "categoria", "precio", "stock", "imagen_url", "grupo"])
        df_prods['grupo_final'] = df_prods.apply(lambda r: r['grupo'] if (r['grupo'] and r['grupo'] != '📦 Varios y Novedades') else obtener_grupo(r['categoria'], r['nombre']), axis=1)
        
        if busqueda:
            df_filtrado = df_prods[df_prods['nombre'].str.contains(busqueda, case=False, na=False) | 
                                   df_prods['categoria'].str.contains(busqueda, case=False, na=False)]
            st.subheader(f"Resultados de búsqueda ({len(df_filtrado)})")
            grupos_mostrar = {"🔍 Resultados": df_filtrado}
        else:
            grupos_orden = [
                "🔑 Llaveros, Pines y Dijes",
                "🧩 Juguetes y Fidgets",
                "🗿 Figuras y Funkos",
                "🏠 Hogar y Deco",
                "📦 Varios y Novedades"
            ]
            grupos_mostrar = {g: df_prods[df_prods['grupo_final'] == g] for g in grupos_orden if not df_prods[df_prods['grupo_final'] == g].empty}

        for nombre_grupo, df_g in grupos_mostrar.items():
            with st.expander(f"{nombre_grupo} ({len(df_g)} artículos)", expanded=True if busqueda else False):
                cols = st.columns(3)
                for idx, row in df_g.reset_index().iterrows():
                    col = cols[idx % 3]
                    with col:
                        st.markdown("---")
                        img_val = row['imagen_url']
                        if pd.notnull(img_val) and isinstance(img_val, str) and img_val.startswith('data:image'):
                            st.image(img_val, use_container_width=True)
                        else:
                            st.caption("📷 *Sin foto miniatura*")
                        
                        st.markdown(f"**{row['nombre']}**")
                        st.caption(f"Categoría: {row['categoria']}")
                        st.subheader(f"${float(row['precio']):,.0f}")
                        st.write(f"Stock: **{row['stock']} un.**")
                        
                        if row['stock'] > 0:
                            if st.button(f"🛒 Vender 1 un.", key=f"btn_{key_prefix}_{row['id']}", use_container_width=True):
                                subt = float(row['precio'])
                                consulta("UPDATE productos SET stock = stock - 1 WHERE id = %s", (row['id'],), fetch=False)
                                consulta("""
                                    INSERT INTO ventas (origen, nombre_feria, item_tipo, item_nombre, cantidad, precio_unitario, subtotal, fecha) 
                                    VALUES (%s, %s, 'PRODUCTO', %s, 1, %s, %s, %s)
                                """, (origen_venta, nombre_feria, row['nombre'], row['precio'], subt, datetime.now()), fetch=False)
                                st.success(f"Vendido en {origen_venta}: {row['nombre']}")
                                st.rerun()
                        else:
                            st.error("Sin Stock")
    else:
        st.info("Aún no hay productos cargados en el inventario.")

    # SECCIÓN GACHAPON
    if incluir_gachapon:
        with st.expander("🎰 Gachapon", expanded=False):
            premios_all = consulta("SELECT id, numero, nombre, coleccion, stock_deposito, en_maquina, precio_ficha FROM gachapon_premios ORDER BY id ASC")
            
            if premios_all:
                df_gach = pd.DataFrame(premios_all, columns=["ID", "Nº", "Premio", "Colección", "En Depósito", "En Máquina", "Precio Ficha"])
                df_gach['Colección'] = df_gach['Colección'].fillna('General')
                colecciones_unicas = sorted(df_gach['Colección'].unique())
                
                st.subheader(f"🎟️ Venta de Ficha Gachapon ({origen_venta})")
                col_sel = st.selectbox("Seleccionar Colección:", colecciones_unicas, key=f"gach_sel_{key_prefix}")
                
                p_ficha = float(df_gach[df_gach['Colección'] == col_sel]['Precio Ficha'].iloc[0])
                st.caption(f"Precio actual por ficha ({col_sel}): **${p_ficha:,.0f}**")
                
                c_f1, c_f2 = st.columns([2, 1])
                with c_f1:
                    cant_fichas = st.number_input("Cantidad de Fichas:", min_value=1, value=1, key=f"cant_f_{key_prefix}")
                with c_f2:
                    st.write("")
                    if st.button("🎟️ Vender Ficha/s", key=f"btn_f_{key_prefix}", use_container_width=True):
                        subt = cant_fichas * p_ficha
                        consulta("""
                            INSERT INTO ventas (origen, nombre_feria, item_tipo, item_nombre, cantidad, precio_unitario, subtotal, fecha) 
                            VALUES (%s, %s, 'FICHA_GACHAPON', %s, %s, %s, %s, %s)
                        """, (origen_venta, nombre_feria, f"Ficha Gachapon ({col_sel})", cant_fichas, p_ficha, subt, datetime.now()), fetch=False)
                        st.success(f"Vendido en {origen_venta}: {cant_fichas} ficha/s {col_sel} (${subt:,.0f})")

                st.divider()
                
                tabs_colec = st.tabs([f"📦 {col}" for col in colecciones_unicas])
                for index, col_nombre in enumerate(colecciones_unicas):
                    with tabs_colec[index]:
                        sub_df = df_gach[df_gach['Colección'] == col_nombre]
                        st.dataframe(sub_df[["Nº", "Premio", "En Máquina", "En Depósito"]], use_container_width=True)
            else:
                st.caption("No hay premios de Gachapon cargados aún.")

# --- INTERFAZ PRINCIPAL ---
st.title("🏪 Control Termas, Taller & Feria")

tab_termas, tab_taller, tab_feria, tab_caja, tab_stock = st.tabs([
    "🛒 Ventas Termas", 
    "🛠️ Ventas Taller", 
    "🎪 Ventas Feria", 
    "💰 Caja & Cierres", 
    "📦 Inventario"
])

prods_db = consulta("SELECT id, nombre, categoria, precio, stock, imagen_url, grupo FROM productos ORDER BY nombre ASC")

# -----------------------------------------------------------------------------
# 1. VENTAS TERMAS
# -----------------------------------------------------------------------------
with tab_termas:
    st.header("🛒 Ventas Mostrador Termas")
    renderizar_catalogo(prods_db, origen_venta="TERMAS", incluir_gachapon=True, key_prefix="termas")

# -----------------------------------------------------------------------------
# 2. VENTAS TALLER
# -----------------------------------------------------------------------------
with tab_taller:
    st.header("🛠️ Ventas Taller / Envíos Externos")
    renderizar_catalogo(prods_db, origen_venta="TALLER", incluir_gachapon=False, key_prefix="taller")

# -----------------------------------------------------------------------------
# 3. VENTAS FERIA
# -----------------------------------------------------------------------------
with tab_feria:
    st.header("🎪 Ventas Stand / Feria")
    
    col_f_nombre, col_f_cierre = st.columns([3, 2])
    with col_f_nombre:
        nombre_feria_activa = st.text_input("Nombre / Lugar de la Feria Activa:", value="Feria Dolores", key="input_nombre_feria")
    with col_f_cierre:
        st.write("")
        res_feria_act = consulta("SELECT SUM(subtotal) FROM ventas WHERE origen='FERIA' AND nombre_feria = %s AND cerrado = 0", (nombre_feria_activa,))
        tot_feria_act = float(res_feria_act[0][0]) if res_feria_act and res_feria_act[0][0] else 0.0
        
        if st.button(f"🔒 Cerrar Feria '{nombre_feria_activa}' (${tot_feria_act:,.0f})", use_container_width=True):
            if tot_feria_act > 0:
                cnt_f = consulta("SELECT COUNT(*) FROM ventas WHERE origen='FERIA' AND nombre_feria = %s AND cerrado = 0", (nombre_feria_activa,))[0][0]
                min_f = consulta("SELECT MIN(fecha) FROM ventas WHERE origen='FERIA' AND nombre_feria = %s AND cerrado = 0", (nombre_feria_activa,))[0][0]
                
                consulta("""
                    INSERT INTO cierres_caja (tipo_cierre, etiqueta_ciclo, total_recaudado, total_operaciones, fecha_inicio, fecha_cierre)
                    VALUES ('FERIA', %s, %s, %s, %s, %s)
                """, (f"Feria: {nombre_feria_activa}", tot_feria_act, cnt_f, min_f, datetime.now()), fetch=False)
                
                id_nuevo_cierre = consulta("SELECT MAX(id) FROM cierres_caja")[0][0]
                consulta("UPDATE ventas SET cerrado = 1, id_cierre = %s WHERE origen='FERIA' AND nombre_feria = %s AND cerrado = 0", (id_nuevo_cierre, nombre_feria_activa), fetch=False)
                st.success(f"¡Feria '{nombre_feria_activa}' cerrada exitosamente!")
                st.rerun()
            else:
                st.info("No hay ventas abiertas para esta feria.")

    st.divider()
    renderizar_catalogo(prods_db, origen_venta="FERIA", incluir_gachapon=True, key_prefix="feria", nombre_feria=nombre_feria_activa)

# -----------------------------------------------------------------------------
# 4. CAJA & CIERRES DE CICLO (CON SELECTOR DE CAJA)
# -----------------------------------------------------------------------------
with tab_caja:
    st.header("💰 Estado de Caja y Cierre de Ciclos")
    
    with st.expander("🔒 Realizar Cierre de Ciclo Específico", expanded=True):
        caja_a_cerrar = st.selectbox(
            "Seleccionar la Caja que querés cerrar:",
            ["🛒 Solo Termas", "🛠️ Solo Taller", "🏪 Termas + Taller (Ambos)", "🎪 Feria Específica"],
            key="sel_caja_cierre"
        )
        
        # Filtros de consulta según la opción elegida
        if caja_a_cerrar == "🛒 Solo Termas":
            where_clause = "origen = 'TERMAS' AND cerrado = 0"
            params_where = ()
            tipo_db = "TERMAS"
        elif caja_a_cerrar == "🛠️ Solo Taller":
            where_clause = "origen = 'TALLER' AND cerrado = 0"
            params_where = ()
            tipo_db = "TALLER"
        elif caja_a_cerrar == "🏪 Termas + Taller (Ambos)":
            where_clause = "origen IN ('TERMAS', 'TALLER') AND cerrado = 0"
            params_where = ()
            tipo_db = "TERMAS_TALLER"
        else: # Feria Específica
            ferias_abiertas = consulta("SELECT DISTINCT nombre_feria FROM ventas WHERE origen = 'FERIA' AND cerrado = 0")
            lista_fa = [f[0] for f in ferias_abiertas] if ferias_abiertas else ["Feria Dolores"]
            feria_sel_cierre = st.selectbox("Seleccionar Feria a cerrar:", lista_fa, key="sel_feria_especifica")
            where_clause = "origen = 'FERIA' AND nombre_feria = %s AND cerrado = 0"
            params_where = (feria_sel_cierre,)
            tipo_db = f"FERIA: {feria_sel_cierre}"

        res_cierre_sel = consulta(f"SELECT SUM(subtotal), COUNT(*) FROM ventas WHERE {where_clause}", params_where)
        tot_sel = float(res_cierre_sel[0][0]) if res_cierre_sel and res_cierre_sel[0][0] else 0.0
        cnt_sel = int(res_cierre_sel[0][1]) if res_cierre_sel and res_cierre_sel[0][1] else 0

        ca1, ca2 = st.columns([3, 2])
        with ca1:
            etiqueta_ciclo = st.text_input("Nombre / Etiqueta del Ciclo a Cerrar:", value=f"Ciclo {caja_a_cerrar} - {datetime.now().strftime('%B %Y')}")
            st.write(f"Monto Total Acumulado Abierto: **${tot_sel:,.0f}** ({cnt_sel} ventas)")
        with ca2:
            st.write("")
            st.write("")
            if st.button("🔒 Ejecutar Cierre de Ciclo", use_container_width=True):
                if tot_sel > 0:
                    min_f_sel = consulta(f"SELECT MIN(fecha) FROM ventas WHERE {where_clause}", params_where)[0][0]
                    consulta("""
                        INSERT INTO cierres_caja (tipo_cierre, etiqueta_ciclo, total_recaudado, total_operaciones, fecha_inicio, fecha_cierre)
                        VALUES (%s, %s, %s, %s, %s, %s)
                    """, (tipo_db, etiqueta_ciclo, tot_sel, cnt_sel, min_f_sel, datetime.now()), fetch=False)
                    
                    id_c = consulta("SELECT MAX(id) FROM cierres_caja")[0][0]
                    consulta(f"UPDATE ventas SET cerrado = 1, id_cierre = %s WHERE {where_clause}", (id_c,) + params_where, fetch=False)
                    st.success(f"¡Cierre '{etiqueta_ciclo}' realizado con éxito por ${tot_sel:,.0f}!")
                    st.rerun()
                else:
                    st.info(f"No hay ventas abiertas en {caja_a_cerrar} para cerrar.")

    st.divider()

    st.subheader("📅 Consulta de Caja por Rango de Fechas")
    col_f1, col_f2 = st.columns(2)
    with col_f1:
        fecha_desde = st.date_input("Fecha Desde:", value=date(date.today().year, date.today().month, 1))
    with col_f2:
        fecha_hasta = st.date_input("Fecha Hasta:", value=date.today())

    ventas_rango = consulta("""
        SELECT id, fecha, origen, nombre_feria, item_tipo, item_nombre, cantidad, precio_unitario, subtotal
        FROM ventas
        WHERE fecha >= %s AND fecha <= %s
        ORDER BY fecha DESC
    """, (datetime.combine(fecha_desde, datetime.min.time()), datetime.combine(fecha_hasta, datetime.max.time())))

    tot_termas_r = sum(float(v[8]) for v in ventas_rango if v[2] == 'TERMAS') if ventas_rango else 0.0
    tot_taller_r = sum(float(v[8]) for v in ventas_rango if v[2] == 'TALLER') if ventas_rango else 0.0
    tot_feria_r = sum(float(v[8]) for v in ventas_rango if v[2] == 'FERIA') if ventas_rango else 0.0
    tot_comb_r = tot_termas_r + tot_taller_r + tot_feria_r

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("🛒 TERMAS (En Rango)", f"${tot_termas_r:,.0f}")
    m2.metric("🛠️ TALLER (En Rango)", f"${tot_taller_r:,.0f}")
    m3.metric("🎪 FERIA (En Rango)", f"${tot_feria_r:,.0f}")
    m4.metric("💵 TOTAL EN RANGO", f"${tot_comb_r:,.0f}")

    if ventas_rango:
        df_v_rango = pd.DataFrame(ventas_rango, columns=["ID", "Fecha", "Origen", "Feria", "Tipo", "Artículo", "Cantidad", "Precio Unitario ($)", "Subtotal ($)"])
        
        buffer_v = io.BytesIO()
        with pd.ExcelWriter(buffer_v, engine='openpyxl') as writer:
            df_v_rango.to_excel(writer, index=False, sheet_name='Reporte_Ventas')
        
        st.download_button(
            label="📥 Descargar Reporte de Ventas del Rango en Excel (.xlsx)",
            data=buffer_v.getvalue(),
            file_name=f"Reporte_Ventas_{fecha_desde}_al_{fecha_hasta}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
        st.dataframe(df_v_rango, use_container_width=True)

# -----------------------------------------------------------------------------
# 5. INVENTARIO CON EXPORTACIÓN E IMPORTACIÓN CON FOTOS EN EXCEL
# -----------------------------------------------------------------------------
with tab_stock:
    st.header("📦 Inventario")
    
    # 1. ALTAS DIRECCIONADAS MANUALES CON FOTO
    with st.expander("➕ Dar de Alta Nuevo Producto o Premio Gachapon", expanded=False):
        tipo_alta = st.radio("¿Qué querés registrar?", ["Producto General", "Premio de Gachapon"], horizontal=True)
        
        if tipo_alta == "Producto General":
            col_a1, col_a2 = st.columns(2)
            with col_a1:
                nuevo_nombre = st.text_input("Nombre del Producto:")
                nuevo_grupo = st.selectbox("Grupo:", [
                    "🔑 Llaveros, Pines y Dijes",
                    "🧩 Juguetes y Fidgets",
                    "🗿 Figuras y Funkos",
                    "🏠 Hogar y Deco",
                    "📦 Varios y Novedades"
                ])
                nueva_cat = st.text_input("Categoría o Subgrupo (ej: Clubes, Anti-estrés):", value="General")
            with col_a2:
                nuevo_precio = st.number_input("Precio de Venta ($):", min_value=0.0, value=1000.0, key="alta_precio_prod")
                nuevo_stock = st.number_input("Stock Inicial:", min_value=0, value=1, key="alta_stock_prod")
                archivo_foto_alta = st.file_uploader("📷 Foto de Producto (PNG, JPG, JPEG):", type=["png", "jpg", "jpeg", "webp"], key="foto_alta_prod")
            
            b64_img_alta = procesar_archivo_imagen(archivo_foto_alta)
            if b64_img_alta:
                st.image(archivo_foto_alta, width=120, caption="Vista previa de miniatura")

            if st.button("💾 Guardar Nuevo Producto", use_container_width=True):
                if nuevo_nombre:
                    try:
                        consulta("""
                            INSERT INTO productos (nombre, grupo, categoria, subgrupo, precio, stock, imagen_url) 
                            VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """, (nuevo_nombre, nuevo_grupo, nueva_cat, nueva_cat, nuevo_precio, nuevo_stock, b64_img_alta), fetch=False)
                        st.success(f"¡Producto '{nuevo_nombre}' agregado correctamente con su foto!")
                        st.rerun()
                    except Exception as ex:
                        st.error(f"Error al guardar: {ex}")
                else:
                    st.warning("Escribí un nombre de producto.")

        else: # Alta Gachapon
            col_g1, col_g2 = st.columns(2)
            with col_g1:
                g_num = st.text_input("Número o Código de Premio:")
                g_nombre = st.text_input("Nombre del Personaje/Premio:")
                
                colec_existentes = consulta("SELECT DISTINCT coleccion FROM gachapon_premios")
                lista_c = [c[0] for c in colec_existentes if c[0]] if colec_existentes else ["Zooki", "Minecraft"]
                
                g_colec_sel = st.selectbox("Seleccionar Colección Existente:", lista_c + ["➕ Crear nueva colección..."])
                if g_colec_sel == "➕ Crear nueva colección...":
                    g_colec = st.text_input("Nombre de la NUEVA Colección (ej: Pokémon):")
                    g_precio_ficha = st.number_input("Precio de la Ficha para esta colección ($):", min_value=0.0, value=2000.0, key="alta_p_ficha")
                else:
                    g_colec = g_colec_sel
                    g_precio_ficha = 2000.0
            with col_g2:
                g_en_maq = st.number_input("Cantidad Inicial en Máquina:", min_value=0, value=1, key="alta_maq_gach")
                g_en_dep = st.number_input("Cantidad Inicial en Depósito:", min_value=0, value=0, key="alta_dep_gach")
                archivo_foto_gach = st.file_uploader("📷 Foto del Personaje (PNG, JPG, JPEG):", type=["png", "jpg", "jpeg", "webp"], key="foto_alta_gach")

            b64_img_gach = procesar_archivo_imagen(archivo_foto_gach)
            if b64_img_gach:
                st.image(archivo_foto_gach, width=120, caption="Vista previa de miniatura")

            if st.button("💾 Guardar Personaje Gachapon", use_container_width=True):
                if g_nombre and g_colec:
                    try:
                        consulta("""
                            INSERT INTO gachapon_premios (numero, nombre, coleccion, en_maquina, stock_deposito, precio_ficha, imagen_url) 
                            VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """, (g_num, g_nombre, g_colec, g_en_maq, g_en_dep, g_precio_ficha, b64_img_gach), fetch=False)
                        st.success(f"¡Personaje '{g_nombre}' guardado en la colección '{g_colec}' con su foto!")
                        st.rerun()
                    except Exception as ex:
                        st.error(f"Error al guardar: {ex}")
                else:
                    st.warning("Completá el nombre del personaje y la colección.")

    # 2. CAMBIO DE PRECIOS, STOCK Y FOTO
    with st.expander("✏️ Cambiar Precios, Stock y Foto", expanded=False):
        tipo_precio = st.radio("Editar datos de:", ["Producto General", "Ficha/Colección Gachapon"], horizontal=True, key="edit_tipo_radio")
        
        if tipo_precio == "Producto General":
            if prods_db:
                p_edit = st.selectbox("Seleccionar Producto:", [p[1] for p in prods_db], key="edit_sel_producto")
                info_p = [p for p in prods_db if p[1] == p_edit][0]
                val_act = info_p[3]
                stock_act = info_p[4]
                foto_act = info_p[5]
                
                cp1, cp2 = st.columns(2)
                with cp1:
                    p_nuevo = st.number_input("Precio de Venta ($):", value=float(val_act), min_value=0.0, key=f"edit_p_{p_edit}")
                    s_nuevo = st.number_input("Stock Disponible (un.):", value=int(stock_act), min_value=0, key=f"edit_s_{p_edit}")
                with cp2:
                    if pd.notnull(foto_act) and isinstance(foto_act, str) and foto_act.startswith('data:image'):
                        st.image(foto_act, width=100, caption="Foto actual")
                    else:
                        st.caption("📷 *Sin foto asignada*")
                    archivo_foto_edit = st.file_uploader("📷 Nueva Foto (PNG, JPG, JPEG):", type=["png", "jpg", "jpeg", "webp"], key=f"foto_edit_{p_edit}")
                
                b64_foto_nueva = procesar_archivo_imagen(archivo_foto_edit) if archivo_foto_edit is not None else foto_act

                if st.button("💾 Actualizar Producto", use_container_width=True, key="btn_actualizar_prod"):
                    consulta("UPDATE productos SET precio = %s, stock = %s, imagen_url = %s WHERE nombre = %s", (p_nuevo, s_nuevo, b64_foto_nueva, p_edit), fetch=False)
                    st.success(f"¡Producto '{p_edit}' actualizado correctamente!")
                    st.rerun()
        else:
            colec_precios = consulta("SELECT DISTINCT coleccion, precio_ficha FROM gachapon_premios")
            if colec_precios:
                col_e = st.selectbox("Seleccionar Colección:", [c[0] for c in colec_precios], key="edit_sel_coleccion")
                p_f_act = [c[1] for c in colec_precios if c[0] == col_e][0]
                pf_nuevo = st.number_input(f"Nuevo Precio de Ficha para {col_e} ($):", value=float(p_f_act if p_f_act else 2000), min_value=0.0, key=f"edit_pf_{col_e}")
                if st.button("💾 Actualizar Precio de Ficha", use_container_width=True, key="btn_actualizar_ficha"):
                    consulta("UPDATE gachapon_premios SET precio_ficha = %s WHERE coleccion = %s", (pf_nuevo, col_e), fetch=False)
                    st.success(f"¡Precio de ficha para {col_e} actualizado a ${pf_nuevo:,.0f}!")
                    st.rerun()

    # 3. IMPORTAR / EXPORTAR INVENTARIO CON IMÁGENES DENTRO DEL EXCEL
    with st.expander("📥 📤 Importar y Exportar Inventario Completo con Fotos en Excel", expanded=True):
        st.write("Generá una planilla Excel con las fotos de los productos incrustadas en las celdas.")
        
        all_p = consulta("SELECT id, nombre, categoria, subgrupo, precio, stock, grupo, imagen_url FROM productos ORDER BY nombre ASC")
        
        def generar_excel_con_fotos(productos):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Productos_Generales"
            
            headers = ["ID", "Foto", "Producto", "Categoría", "Subgrupo", "Precio Venta ($)", "Stock", "Grupo"]
            ws.append(headers)
            
            ws.column_dimensions['B'].width = 15
            
            for idx, p in enumerate(productos, start=2):
                p_id, p_nom, p_cat, p_sub, p_prec, p_stk, p_grp, p_img = p
                ws.row_dimensions[idx].height = 55
                
                ws.cell(row=idx, column=1, value=p_id)
                ws.cell(row=idx, column=3, value=p_nom)
                ws.cell(row=idx, column=4, value=p_cat)
                ws.cell(row=idx, column=5, value=p_sub)
                ws.cell(row=idx, column=6, value=float(p_prec))
                ws.cell(row=idx, column=7, value=int(p_stk))
                ws.cell(row=idx, column=8, value=p_grp)
                
                if pd.notnull(p_img) and isinstance(p_img, str) and "base64," in p_img:
                    try:
                        header, encoded = p_img.split("base64,")
                        img_bytes = base64.b64decode(encoded)
                        pil_img = PILImage.open(io.BytesIO(img_bytes))
                        pil_img.thumbnail((65, 65))
                        
                        img_buf = io.BytesIO()
                        pil_img.save(img_buf, format="PNG")
                        img_buf.seek(0)
                        
                        xl_img = OpenPyXLEImage(img_buf)
                        ws.add_image(xl_img, f"B{idx}")
                    except Exception:
                        pass
                        
            buf = io.BytesIO()
            wb.save(buf)
            return buf.getvalue()

        if all_p:
            excel_con_fotos_bytes = generar_excel_con_fotos(all_p)
            st.download_button(
                label="📤 Exportar Inventario con Fotos a Excel (.xlsx)",
                data=excel_con_fotos_bytes,
                file_name=f"Inventario_Visual_{date.today()}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )

        st.divider()
        
        # IMPORTADOR
        archivo_excel = st.file_uploader("📥 Seleccionar planilla Excel para Sincronizar (.xlsx o .xls):", type=["xlsx", "xls"], key="excel_uploader")
        
        if archivo_excel is not None:
            if st.button("🚀 Sincronizar Base de Datos con Excel", use_container_width=True):
                try:
                    xls = pd.ExcelFile(archivo_excel)
                    cargados_prods = 0
                    cargados_gach = 0
                    
                    for sheet in xls.sheet_names:
                        df_sheet = pd.read_excel(archivo_excel, sheet_name=sheet)
                        df_sheet.columns = [str(c).strip().lower() for c in df_sheet.columns]
                        
                        is_gachapon = any(k in sheet.lower() for k in ['zooki', 'minecraft', 'gachapon', 'premio']) or \
                                      any('impresos' in c or 'deposito' in c or 'maquina' in c for c in df_sheet.columns)
                        
                        if is_gachapon:
                            col_nom = next((c for c in df_sheet.columns if 'nombre' in c or 'premio' in c or 'personaje' in c), None)
                            col_num = next((c for c in df_sheet.columns if 'numero' in c or 'nº' in c or 'num' in c or 'id' in c), None)
                            col_dep = next((c for c in df_sheet.columns if 'deposito' in c or 'depósito' in c or 'stock' in c), None)
                            col_maq = next((c for c in df_sheet.columns if 'maquina' in c or 'máquina' in c), None)
                            
                            if col_nom:
                                for _, r in df_sheet.dropna(subset=[col_nom]).iterrows():
                                    g_nom = str(r[col_nom]).strip()
                                    g_num = str(r[col_num]).strip() if (col_num and pd.notnull(r[col_num])) else ""
                                    g_dep = int(r[col_dep]) if (col_dep and pd.notnull(r[col_dep])) else 0
                                    g_maq = int(r[col_maq]) if (col_maq and pd.notnull(r[col_maq])) else 0
                                    
                                    consulta("""
                                        INSERT INTO gachapon_premios (numero, nombre, coleccion, stock_deposito, en_maquina)
                                        VALUES (%s, %s, %s, %s, %s)
                                        ON CONFLICT (nombre) DO UPDATE SET
                                            numero = EXCLUDED.numero,
                                            coleccion = EXCLUDED.coleccion,
                                            stock_deposito = EXCLUDED.stock_deposito,
                                            en_maquina = EXCLUDED.en_maquina;
                                    """, (g_num, g_nom, sheet, g_dep, g_maq), fetch=False)
                                    cargados_gach += 1
                        else:
                            col_nombre = next((c for c in df_sheet.columns if 'producto' in c or 'nombre' in c or 'item' in c), None)
                            col_precio = next((c for c in df_sheet.columns if 'precio' in c or 'valor' in c), None)
                            col_stock = next((c for c in df_sheet.columns if 'stock' in c or 'cantidad' in c or 'cant' in c), None)
                            col_cat = next((c for c in df_sheet.columns if 'categoria' in c or 'categoría' in c or 'grupo' in c), None)
                            
                            if col_nombre:
                                for _, r in df_sheet.dropna(subset=[col_nombre]).iterrows():
                                    p_nombre = str(r[col_nombre]).strip()
                                    p_precio = float(r[col_precio]) if (col_precio and pd.notnull(r[col_precio])) else 0.0
                                    p_stock = int(r[col_stock]) if (col_stock and pd.notnull(r[col_stock])) else 0
                                    p_cat = str(r[col_cat]).strip() if (col_cat and pd.notnull(r[col_cat])) else "General"
                                    p_grupo = obtener_grupo(p_cat, p_nombre)
                                    
                                    consulta("""
                                        INSERT INTO productos (nombre, categoria, grupo, precio, stock)
                                        VALUES (%s, %s, %s, %s, %s)
                                        ON CONFLICT (nombre) DO UPDATE SET
                                            precio = EXCLUDED.precio,
                                            stock = EXCLUDED.stock,
                                            categoria = EXCLUDED.categoria,
                                            grupo = EXCLUDED.grupo;
                                    """, (p_nombre, p_cat, p_grupo, p_precio, p_stock), fetch=False)
                                    cargados_prods += 1
                                
                    st.success(f"¡Éxito! Se sincronizaron {cargados_prods} productos generales y {cargados_gach} premios de Gachapon.")
                    st.rerun()
                except Exception as ex:
                    st.error(f"Error al procesar el archivo Excel: {ex}")

    # 4. ARQUEO DE MAQUINA GACHAPON
    with st.expander("🔍 Arqueo y Recompuesto de Máquina Gachapon", expanded=False):
        premios_arq = consulta("SELECT id, numero, nombre, coleccion, en_maquina, stock_deposito FROM gachapon_premios ORDER BY coleccion, numero ASC")
        if premios_arq:
            df_arq = pd.DataFrame(premios_arq, columns=["id", "numero", "nombre", "coleccion", "en_maquina", "stock_deposito"])
            col_arq_sel = st.selectbox("Seleccionar Colección para Auditar/Arquear:", df_arq['coleccion'].unique(), key="arq_col")
            
            df_col_arq = df_arq[df_arq['coleccion'] == col_arq_sel]
            st.write("Ajustá la **Cantidad Real Actual en Máquina** observada al abrir la máquina:")
            
            for index, row in df_col_arq.iterrows():
                ca1, ca2 = st.columns([3, 2])
                with ca1:
                    st.write(f"**Nº {row['numero']} - {row['nombre']}** (En dep: {row['stock_deposito']})")
                with ca2:
                    n_cant_maq = st.number_input(f"En Máquina:", min_value=0, value=int(row['en_maquina']), key=f"arq_n_{row['id']}")
                    if n_cant_maq != int(row['en_maquina']):
                        consulta("UPDATE gachapon_premios SET en_maquina = %s WHERE id = %s", (n_cant_maq, row['id']), fetch=False)
            st.caption("Los cambios de stock se guardan en tiempo real al modificar los números.")

    st.divider()
    st.subheader("📋 Lista Completa de Stock General")
    prods_full = consulta("SELECT id, nombre, grupo, categoria, precio, stock FROM productos ORDER BY nombre ASC")
    if prods_full:
        df_stock = pd.DataFrame(prods_full, columns=["ID", "Producto", "Grupo", "Categoría", "Precio Venta", "Stock"])
        st.dataframe(df_stock, use_container_width=True)
